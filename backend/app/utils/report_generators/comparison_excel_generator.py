from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.report import ComparisonReportDTO


GREEN = "0B6B2B"
LIGHT_GREEN = "EAF6EC"
BORDER_COLOR = "DDE7DF"
TEXT_COLOR = "0F172A"
MUTED_COLOR = "526174"


def _format_datetime(value) -> str:
    if value is None:
        return "Not available"

    return value.strftime("%d %b %Y, %H:%M:%S")


def _format_optional(value) -> str:
    if value is None or value == "":
        return "Not available"

    return str(value)


def _format_boolean(value) -> str:
    if value is True:
        return "Yes"

    if value is False:
        return "No"

    return "Not available"


def generate_comparison_excel(
    report: ComparisonReportDTO,
) -> bytes:
    """
    Generate an XLSX comparison report from stored results.

    No benchmark execution is performed.
    """

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Comparison"

    # ---------------------------------------------------------
    # Styles
    # ---------------------------------------------------------

    title_font = Font(
        name="Calibri",
        size=18,
        bold=True,
        color="FFFFFF",
    )

    subtitle_font = Font(
        name="Calibri",
        size=11,
        italic=True,
        color=MUTED_COLOR,
    )

    section_font = Font(
        name="Calibri",
        size=11,
        bold=True,
        color=GREEN,
    )

    label_font = Font(
        name="Calibri",
        size=10,
        bold=True,
        color=TEXT_COLOR,
    )

    normal_font = Font(
        name="Calibri",
        size=10,
        color=TEXT_COLOR,
    )

    header_font = Font(
        name="Calibri",
        size=10,
        bold=True,
        color="FFFFFF",
    )

    green_fill = PatternFill(
        "solid",
        fgColor=GREEN,
    )

    light_green_fill = PatternFill(
        "solid",
        fgColor=LIGHT_GREEN,
    )

    white_fill = PatternFill(
        "solid",
        fgColor="FFFFFF",
    )

    thin_side = Side(
        style="thin",
        color=BORDER_COLOR,
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    # ---------------------------------------------------------
    # Title
    # ---------------------------------------------------------

    worksheet.merge_cells("A1:G1")
    worksheet["A1"] = "GREEN CODE ANALYZER"
    worksheet["A1"].font = title_font
    worksheet["A1"].fill = green_fill
    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    worksheet.row_dimensions[1].height = 28

    worksheet.merge_cells("A2:G2")
    worksheet["A2"] = "Implementation Comparison Report"
    worksheet["A2"].font = subtitle_font
    worksheet["A2"].alignment = Alignment(
        horizontal="center",
    )

    # ---------------------------------------------------------
    # Comparison Information
    # ---------------------------------------------------------

    row = 4

    worksheet.cell(row=row, column=1).value = (
        "Comparison Information"
    )
    worksheet.cell(row=row, column=1).font = section_font

    row += 1

    comparison_fields = [
        ("Comparison ID", report.comparison_id),
        ("Benchmark Task", report.benchmark_name),
        ("Category", _format_optional(report.category)),
        ("Benchmark Size", _format_optional(report.bench_size)),
        ("Input Size", _format_optional(report.input_size)),
        ("Workload Type", _format_optional(report.workload_type)),
        ("Created At", _format_datetime(report.created_at)),
        ("Implementations", len(report.results)),
    ]

    for label, value in comparison_fields:
        worksheet.cell(row=row, column=1).value = label
        worksheet.cell(row=row, column=1).font = label_font
        worksheet.cell(row=row, column=1).fill = light_green_fill
        worksheet.cell(row=row, column=1).border = border

        worksheet.cell(row=row, column=2).value = value
        worksheet.cell(row=row, column=2).font = normal_font
        worksheet.cell(row=row, column=2).border = border

        row += 1

    # ---------------------------------------------------------
    # Implementation Results
    # ---------------------------------------------------------

    row += 1

    worksheet.cell(row=row, column=1).value = (
        "Implementation Results"
    )
    worksheet.cell(row=row, column=1).font = section_font

    row += 1

    headers = [
        "Language",
        "Compiler",
        "Version",
        "Analysis ID",
        "Execution Time (ms)",
        "CPU Usage (%)",
        "Memory Usage (MB)",
        "Energy Consumption (J)",
        "Green Score",
        "Green Score Label",
        "Output Verified",
    ]

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=row,
            column=column,
        )
        cell.value = header
        cell.font = header_font
        cell.fill = green_fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    header_row = row
    row += 1

    for result in report.results:
        values = [
            result.language,
            result.compiler or "Not available",
            result.version or "Not available",
            result.analysis_id,
            (
                round(result.execution_time, 2)
                if result.execution_time is not None
                else "Not available"
            ),
            (
                round(result.cpu_usage, 2)
                if result.cpu_usage is not None
                else "Not available"
            ),
            (
                round(result.memory_usage, 2)
                if result.memory_usage is not None
                else "Not available"
            ),
            (
                round(result.energy_consumption, 4)
                if result.energy_consumption is not None
                else "Not available"
            ),
            (
                round(result.green_score, 2)
                if result.green_score is not None
                else "Not available"
            ),
            result.green_score_label or "Not available",
            _format_boolean(result.output_verified),
        ]

        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(
                row=row,
                column=column,
            )
            cell.value = value
            cell.font = normal_font
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if row % 2 == 0:
                cell.fill = white_fill

        row += 1

    # ---------------------------------------------------------
    # Validation
    # ---------------------------------------------------------

    row += 1

    worksheet.cell(row=row, column=1).value = "Validation"
    worksheet.cell(row=row, column=1).font = section_font

    row += 1

    validation_headers = [
        "Language",
        "Analysis ID",
        "Input Size",
        "Benchmark Size",
        "Output Verified",
    ]

    for column, header in enumerate(
        validation_headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=row,
            column=column,
        )
        cell.value = header
        cell.font = header_font
        cell.fill = green_fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
        )

    row += 1

    for result in report.results:
        values = [
            result.language,
            result.analysis_id,
            _format_optional(result.input_size),
            _format_optional(result.bench_size),
            _format_boolean(result.output_verified),
        ]

        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(
                row=row,
                column=column,
            )
            cell.value = value
            cell.font = normal_font
            cell.border = border

        row += 1

    # ---------------------------------------------------------
    # Footer note
    # ---------------------------------------------------------

    row += 2

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=11,
    )

    worksheet.cell(row=row, column=1).value = (
        "This comparison report is generated from stored "
        "benchmark analysis results. Generating the report "
        "does not rerun the benchmark."
    )

    worksheet.cell(row=row, column=1).font = Font(
        name="Calibri",
        size=9,
        italic=True,
        color=MUTED_COLOR,
    )

    worksheet.cell(row=row, column=1).alignment = Alignment(
        wrap_text=True,
    )

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=11,
    )

    worksheet.cell(row=row, column=1).value = (
        "Energy values represent measured energy consumption "
        "during the stored benchmark execution."
    )

    worksheet.cell(row=row, column=1).font = Font(
        name="Calibri",
        size=9,
        italic=True,
        color=MUTED_COLOR,
    )

    worksheet.cell(row=row, column=1).alignment = Alignment(
        wrap_text=True,
    )

    # ---------------------------------------------------------
    # Formatting
    # ---------------------------------------------------------

    worksheet.freeze_panes = f"A{header_row + 1}"

    column_widths = {
        "A": 18,
        "B": 18,
        "C": 14,
        "D": 14,
        "E": 20,
        "F": 17,
        "G": 20,
        "H": 23,
        "I": 15,
        "J": 20,
        "K": 18,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # ---------------------------------------------------------
    # Save
    # ---------------------------------------------------------

    buffer = BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()