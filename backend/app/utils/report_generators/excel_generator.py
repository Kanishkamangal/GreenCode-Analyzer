from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.schemas.report import ReportDTO


def _format_datetime(value) -> str:
    return value.strftime(
        "%d %b %Y, %H:%M:%S"
    )


def _format_optional(value) -> str:
    if value is None or value == "":
        return "Not available"

    return str(value)


def _format_boolean(value) -> str:
    if value is True:
        return "Yes"

    if value is False:
        return "No"

    return "Not available"


def generate_excel(
    report: ReportDTO,
) -> bytes:

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Report"

    # ---------------------------------------------------------
    # STYLES
    # ---------------------------------------------------------

    title_font = Font(
        bold=True,
        size=18,
    )

    subtitle_font = Font(
        italic=True,
        size=11,
    )

    section_font = Font(
        bold=True,
        size=12,
    )

    label_font = Font(
        bold=True,
    )

    section_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    label_fill = PatternFill(
        fill_type="solid",
        fgColor="F3F4F6",
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # ---------------------------------------------------------
    # COLUMN WIDTHS
    # ---------------------------------------------------------

    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 45

    # ---------------------------------------------------------
    # HEADER
    # ---------------------------------------------------------

    worksheet.merge_cells(
        "A1:B1"
    )

    worksheet["A1"] = "GREEN CODE ANALYZER"
    worksheet["A1"].font = title_font
    worksheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    worksheet.merge_cells(
        "A2:B2"
    )

    worksheet["A2"] = (
        "Performance & Sustainability Report"
    )

    worksheet["A2"].font = subtitle_font
    worksheet["A2"].alignment = Alignment(
        horizontal="center"
    )

    worksheet.merge_cells(
        "A4:B4"
    )

    worksheet["A4"] = report.report_title
    worksheet["A4"].font = Font(
        bold=True,
        size=14,
    )

    # ---------------------------------------------------------
    # BENCHMARK INFORMATION
    # ---------------------------------------------------------

    row = 6

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    worksheet.cell(
        row=row,
        column=1,
        value="Benchmark Information",
    )

    worksheet.cell(
        row=row,
        column=1,
    ).font = section_font

    worksheet.cell(
        row=row,
        column=1,
    ).fill = section_fill

    row += 1

    benchmark_data = [
        (
            "Benchmark Task",
            report.metadata.benchmark_task,
        ),
        (
            "Category",
            _format_optional(
                report.metadata.category
            ),
        ),
        (
            "Analysis Type",
            report.metadata.analysis_type.title(),
        ),
        (
            "Language",
            report.metadata.language,
        ),
        (
            "Input Size",
            _format_optional(
                report.metadata.input_size
            ),
        ),
        (
            "Benchmark Size",
            _format_optional(
                report.metadata.bench_size
            ),
        ),
        (
            "Workload Type",
            _format_optional(
                report.metadata.workload_type
            ),
        ),
        (
            "Execution Date",
            _format_datetime(
                report.metadata.execution_date
            ),
        ),
    ]

    row = _write_section(
        worksheet,
        row,
        benchmark_data,
        label_font,
        label_fill,
        thin_border,
    )

    # ---------------------------------------------------------
    # EXECUTION ENVIRONMENT
    # ---------------------------------------------------------

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    worksheet.cell(
        row=row,
        column=1,
        value="Execution Environment",
    )

    worksheet.cell(
        row=row,
        column=1,
    ).font = section_font

    worksheet.cell(
        row=row,
        column=1,
    ).fill = section_fill

    row += 1

    environment_data = [
        (
            "Runtime",
            _format_optional(
                report.metadata.runtime
            ),
        ),
        (
            "Compiler",
            _format_optional(
                report.metadata.compiler
            ),
        ),
    ]

    row = _write_section(
        worksheet,
        row,
        environment_data,
        label_font,
        label_fill,
        thin_border,
    )

    # ---------------------------------------------------------
    # PERFORMANCE METRICS
    # ---------------------------------------------------------

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    worksheet.cell(
        row=row,
        column=1,
        value="Performance Metrics",
    )

    worksheet.cell(
        row=row,
        column=1,
    ).font = section_font

    worksheet.cell(
        row=row,
        column=1,
    ).fill = section_fill

    row += 1

    performance_data = [
        (
            "Execution Time",
            f"{report.metrics.execution_time:,.2f} ms",
        ),
        (
            "CPU Usage",
            f"{report.metrics.cpu_usage:,.2f} %",
        ),
        (
            "Memory Usage",
            f"{report.metrics.memory_usage:,.2f} MB",
        ),
    ]

    row = _write_section(
        worksheet,
        row,
        performance_data,
        label_font,
        label_fill,
        thin_border,
    )

    # ---------------------------------------------------------
    # SUSTAINABILITY METRICS
    # ---------------------------------------------------------

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    worksheet.cell(
        row=row,
        column=1,
        value="Sustainability Metrics",
    )

    worksheet.cell(
        row=row,
        column=1,
    ).font = section_font

    worksheet.cell(
        row=row,
        column=1,
    ).fill = section_fill

    row += 1

    energy_value = (
        f"{report.metrics.energy_consumption:,.4f} J"
        if report.metrics.energy_consumption is not None
        else "Not available"
    )

    green_score_value = (
        f"{report.metrics.green_score:,.2f} / 100"
        if report.metrics.green_score is not None
        else "Not available"
    )

    sustainability_data = [
        (
            "Energy Consumption",
            energy_value,
        ),
        (
            "Green Score",
            green_score_value,
        ),
        (
            "Green Score Label",
            _format_optional(
                report.metrics.green_score_label
            ),
        ),
    ]

    row = _write_section(
        worksheet,
        row,
        sustainability_data,
        label_font,
        label_fill,
        thin_border,
    )

    # ---------------------------------------------------------
    # VALIDATION
    # ---------------------------------------------------------

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    worksheet.cell(
        row=row,
        column=1,
        value="Validation",
    )

    worksheet.cell(
        row=row,
        column=1,
    ).font = section_font

    worksheet.cell(
        row=row,
        column=1,
    ).fill = section_fill

    row += 1

    validation_data = [
        (
            "Output Verified",
            _format_boolean(
                report.metrics.output_verified
            ),
        ),
    ]

    row = _write_section(
        worksheet,
        row,
        validation_data,
        label_font,
        label_fill,
        thin_border,
    )

    # ---------------------------------------------------------
    # COMPARISON INFORMATION
    # ---------------------------------------------------------

    if report.rank is not None:

        row += 1

        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=2,
        )

        worksheet.cell(
            row=row,
            column=1,
            value="Comparison Information",
        )

        worksheet.cell(
            row=row,
            column=1,
        ).font = section_font

        worksheet.cell(
            row=row,
            column=1,
        ).fill = section_fill

        row += 1

        comparison_data = [
            (
                "Rank",
                str(report.rank),
            ),
        ]

        if report.total_comparisons is not None:
            comparison_data.append(
                (
                    "Total Comparisons",
                    str(report.total_comparisons),
                )
            )

        row = _write_section(
            worksheet,
            row,
            comparison_data,
            label_font,
            label_fill,
            thin_border,
        )

    # ---------------------------------------------------------
    # FOOTER
    # ---------------------------------------------------------

    row += 2

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    worksheet.cell(
        row=row,
        column=1,
        value=(
            "This report is generated from stored "
            "benchmark analysis results. Generating "
            "the report does not rerun the benchmark."
        ),
    )

    worksheet.cell(
        row=row,
        column=1,
    ).alignment = Alignment(
        wrap_text=True
    )

    # ---------------------------------------------------------
    # FREEZE PANES
    # ---------------------------------------------------------

    worksheet.freeze_panes = "A6"

    # ---------------------------------------------------------
    # SAVE TO MEMORY
    # ---------------------------------------------------------

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()


def _write_section(
    worksheet,
    start_row,
    data,
    label_font,
    label_fill,
    thin_border,
):

    row = start_row

    for label, value in data:

        label_cell = worksheet.cell(
            row=row,
            column=1,
            value=label,
        )

        value_cell = worksheet.cell(
            row=row,
            column=2,
            value=value,
        )

        label_cell.font = label_font
        label_cell.fill = label_fill

        label_cell.border = thin_border
        value_cell.border = thin_border

        label_cell.alignment = Alignment(
            vertical="center"
        )

        value_cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

        row += 1

    return row